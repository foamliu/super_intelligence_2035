#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
将《超级智能2035》新版本第一部分（1-6章）合并成一个Word文档
"""

import os
import re
from pathlib import Path

# 需要安装: pip install python-docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn

def set_chinese_font(run, font_name='SimSun', font_size=12, bold=False, italic=False):
    """设置中文字体"""
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    # 设置中文字体（关键）
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)

def get_new_chapter_files():
    """获取chapters_new目录下的所有章节文件，按正确顺序排序"""
    chapters_dir = Path("chapters_new")
    files = list(chapters_dir.glob("*.md"))
    
    # 按文件名排序（01-, 02-, 03-...）
    def sort_key(f):
        # 提取数字前缀
        match = re.match(r'(\d+)-', f.name)
        if match:
            return int(match.group(1))
        return 999  # 其他文件放最后
    
    sorted_files = sorted(files, key=sort_key)
    
    # 只取前6章（第一部分）
    part1_files = []
    for f in sorted_files:
        match = re.match(r'(\d+)', f.name)
        if match and int(match.group(1)) <= 6:
            part1_files.append(f)
    
    return part1_files

def add_heading(doc, text, level=1):
    """添加标题，并设置为Word标题样式（用于目录）"""
    # Word的标题级别：Heading 1, Heading 2, etc.
    style_name = f'Heading {level}'
    
    # 确保样式存在并设置中文字体
    try:
        style = doc.styles[style_name]
    except KeyError:
        style = doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
    
    # 设置样式字体
    style.font.name = 'SimHei'  # 黑体用于标题
    style.font.size = Pt([18, 16, 14, 12][min(level-1, 3)])
    style.font.bold = True
    
    # 关键：设置中文字体
    style.element.rPr.rFonts.set(qn('w:eastAsia'), 'SimHei')
    
    paragraph = doc.add_paragraph(text, style=style_name)
    
    # 为每个run设置中文字体
    for run in paragraph.runs:
        set_chinese_font(run, 'SimHei', [18, 16, 14, 12][min(level-1, 3)], bold=True)
    
    return paragraph

def process_markdown_content(doc, content):
    """处理markdown内容并添加到文档（简化版本）"""
    lines = content.split('\n')
    
    for line in lines:
        line = line.rstrip()
        
        # 标题
        if line.startswith('# '):
            add_heading(doc, line[2:].strip(), 1)
        elif line.startswith('## '):
            add_heading(doc, line[3:].strip(), 2)
        elif line.startswith('### '):
            add_heading(doc, line[4:].strip(), 3)
        elif line.startswith('#### '):
            add_heading(doc, line[5:].strip(), 4)
        
        # 列表项
        elif line.startswith('- ') or line.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            run = p.add_run(line[2:].strip())
            set_chinese_font(run, 'SimSun', 12)
        
        # 编号列表
        elif re.match(r'^\d+\.\s', line):
            p = doc.add_paragraph(style='List Number')
            text = re.sub(r'^\d+\.\s', '', line)
            run = p.add_run(text)
            set_chinese_font(run, 'SimSun', 12)
        
        # 图片（简化为文本）
        elif line.startswith('!['):
            # 提取图片描述
            match = re.match(r'!\[([^\]]*)\]\(([^)]+)\)', line)
            if match:
                alt_text, img_path = match.groups()
                p = doc.add_paragraph()
                run = p.add_run(f'[图片: {alt_text}]')
                set_chinese_font(run, 'SimSun', 10, italic=True)
        
        # 图片说明（斜体）
        elif line.startswith('*') and line.endswith('*'):
            p = doc.add_paragraph()
            run = p.add_run(line[1:-1])
            set_chinese_font(run, 'SimSun', 10, italic=True)
        
        # 引用
        elif line.startswith('> '):
            p = doc.add_paragraph(style='Quote')
            run = p.add_run(line[2:])
            set_chinese_font(run, 'SimSun', 12)
        
        # 分隔线
        elif line.startswith('---'):
            # 添加空行作为分隔
            doc.add_paragraph()
        
        # 普通段落（非空行）
        elif line.strip():
            p = doc.add_paragraph()
            run = p.add_run(line)
            set_chinese_font(run, 'SimSun', 12)
        
        # 空行
        else:
            doc.add_paragraph()

def setup_document_styles(doc):
    """设置文档样式"""
    # 设置默认样式
    style = doc.styles['Normal']
    style.font.name = 'SimSun'
    style.font.size = Pt(12)
    # 关键：设置中文字体
    style.element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')
    
    # 设置列表样式
    for style_name in ['List Bullet', 'List Number']:
        try:
            style = doc.styles[style_name]
            style.font.name = 'SimSun'
            style.element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')
        except:
            pass

def add_page_numbers(doc):
    """为文档添加页码"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn as oqn
    
    # 为每个节添加页脚
    for section in doc.sections:
        footer = section.footer
        footer_para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # 添加页码字段
        run = footer_para.add_run()
        
        # 创建fldChar元素
        fldChar1 = OxmlElement('w:fldChar')
        fldChar1.set(oqn('w:fldCharType'), 'begin')
        
        instrText = OxmlElement('w:instrText')
        instrText.set(oqn('xml:space'), 'preserve')
        instrText.text = "PAGE"
        
        fldChar2 = OxmlElement('w:fldChar')
        fldChar2.set(oqn('w:fldCharType'), 'end')
        
        run._r.append(fldChar1)
        run._r.append(instrText)
        run._r.append(fldChar2)
        
        # 设置页码字体
        set_chinese_font(run, 'SimSun', 10)

def main():
    print("开始合并《超级智能2035》新版本第一部分（1-6章）...")
    
    # 获取章节文件
    chapter_files = get_new_chapter_files()
    print(f"找到 {len(chapter_files)} 个章节文件（第一部分）")
    
    if not chapter_files:
        print("错误：未找到章节文件！")
        return
    
    # 创建Word文档
    doc = Document()
    
    # 设置文档样式
    setup_document_styles(doc)
    
    # 添加封面
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('超级智能2035')
    set_chinese_font(run, 'SimHei', 28, bold=True)
    
    doc.add_paragraph()
    
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run('第一部分：身在此山中（1-6章）')
    set_chinese_font(run, 'SimHei', 18, bold=True)
    
    doc.add_paragraph()
    
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = info.add_run('事缓则圆——在加速时代的生存哲学')
    set_chinese_font(run, 'SimHei', 14, bold=True)
    
    doc.add_page_break()
    
    # 添加目录标题
    toc_title = doc.add_paragraph()
    toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = toc_title.add_run('目  录')
    set_chinese_font(run, 'SimHei', 18, bold=True)
    
    doc.add_paragraph()
    
    # 添加章节列表
    for i, file_path in enumerate(chapter_files, 1):
        with open(file_path, 'r', encoding='utf-8') as f:
            first_line = f.readline().strip()
            title = first_line[2:] if first_line.startswith('# ') else first_line
            
        p = doc.add_paragraph(style='List Number')
        run = p.add_run(f"第{i}章: {title}")
        set_chinese_font(run, 'SimSun', 12)
    
    doc.add_page_break()
    
    # 处理每个章节
    for i, file_path in enumerate(chapter_files, 1):
        print(f"处理第 {i}/{len(chapter_files)} 个文件: {file_path.name}")
        
        # 读取markdown内容
        with open(file_path, 'r', encoding='utf-8') as f:
            content = f.read()
        
        # 添加分页符（除了第一章）
        if i > 1:
            doc.add_page_break()
        
        # 处理内容
        process_markdown_content(doc, content)
    
    # 添加页码
    print("添加页码...")
    add_page_numbers(doc)
    
    # 保存文档
    output_file = '超级智能2035_第一部分_身在此山中.docx'
    doc.save(output_file)
    print(f"\n完成！文档已保存为: {output_file}")
    print(f"包含章节: {len(chapter_files)}章")
    print("提示：请在Word中打开文档查看完整格式")

if __name__ == '__main__':
    main()